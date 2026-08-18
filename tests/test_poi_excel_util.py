from openpyxl import Workbook

from script_py.util.PoiExcelUtil import PoiExcelUtil


def test_merged_header_write_read(tmp_path):
    out = tmp_path / "merged.xlsx"
    wb = Workbook()
    sheet = wb.active.title

    # 两列、两行表头（各自纵向合并）
    lstHeader = [["序号", "序号"], ["名称", "名称"]]
    PoiExcelUtil.writeHeaderByCol(wb, sheet, lstHeader, 1)

    # 数据从第 2 行开始；第 2 行两列合并，右侧为空
    PoiExcelUtil.writeCellData(wb, sheet, 2, 0, 1)
    PoiExcelUtil.writeCellData(wb, sheet, 2, 1, "")
    wb[sheet].merge_cells("A3:B3")
    for i in range(1, 5):
        PoiExcelUtil.writeCellData(wb, sheet, 2 + i, 0, i + 1)
        PoiExcelUtil.writeCellData(wb, sheet, 2 + i, 1, "商品%d" % i)
    PoiExcelUtil.write(wb, str(out))

    data = PoiExcelUtil.toMapFull(str(out), sheet, 0, 0, -1, 2, -1, None)
    assert len(data) == 5
    assert data[0]["序号"] == "1"
    # 合并单元格：第 2 行第 1 列（名称）为空，应取合并区左上角值
    assert data[0]["名称"] == "1"
    assert data[1]["名称"] == "商品1"


def test_explicit_style_write_read(tmp_path):
    out = tmp_path / "style.xlsx"
    wb = Workbook()
    sheet = wb.active.title

    lstHeader = [["序号"], ["名称"]]
    styles = PoiExcelUtil.getCellStyles(wb)
    PoiExcelUtil.writeHeaderByStyle(wb, sheet, lstHeader, styles, 1)

    rowIndex = 1
    for i in range(5):
        PoiExcelUtil.writeCellDataByStyle(wb, sheet, styles, rowIndex, 0, i + 1)
        PoiExcelUtil.writeCellDataByStyle(wb, sheet, styles, rowIndex, 1, "商品%d" % i)
        rowIndex += 1
    PoiExcelUtil.write(wb, str(out))

    data = PoiExcelUtil.toMap(str(out))
    assert len(data) == 5
    assert data[2]["序号"] == "3"
    assert data[4]["名称"] == "商品4"


def test_state_isolation_across_files(tmp_path):
    lstHeader = [["序号"], ["名称"]]

    out1 = tmp_path / "a.xlsx"
    wb1 = Workbook()
    PoiExcelUtil.writeHeaderByCol(wb1, wb1.active.title, lstHeader, 1)
    for i in range(5):
        PoiExcelUtil.writeRowData(wb1, wb1.active.title, {0: i + 1, 1: "A-%d" % i})
    PoiExcelUtil.write(wb1, str(out1))

    out2 = tmp_path / "b.xlsx"
    wb2 = Workbook()
    PoiExcelUtil.writeHeaderByCol(wb2, wb2.active.title, lstHeader, 1)
    for i in range(5):
        PoiExcelUtil.writeRowData(wb2, wb2.active.title, {0: i + 1, 1: "B-%d" % i})
    PoiExcelUtil.write(wb2, str(out2))

    data1 = PoiExcelUtil.toMap(str(out1))
    data2 = PoiExcelUtil.toMap(str(out2))
    assert len(data1) == 5 and len(data2) == 5
    assert data1[0]["名称"] == "A-0"
    assert data2[0]["名称"] == "B-0"


def test_write_header_by_width_col(tmp_path):
    out = tmp_path / "widthcol.xlsx"
    wb = Workbook()
    sheet = wb.active.title
    styles = PoiExcelUtil.getCellStyles(wb)
    lstHeader = [["序号"], ["名称"]]
    PoiExcelUtil.writeHeaderByWidthCol(wb, sheet, lstHeader, styles, 1, 18)
    PoiExcelUtil.writeCellDataByStyle(wb, sheet, styles, 1, 0, 1)
    PoiExcelUtil.writeCellDataByStyle(wb, sheet, styles, 1, 1, "商品")
    PoiExcelUtil.write(wb, str(out))

    # 固定列宽 18 作为下限（短数据自动列宽 13 < 18，最终为 18）
    width = wb[sheet].column_dimensions["A"].width
    assert width == 18.0


def test_write_header_exclude_row(tmp_path):
    out = tmp_path / "exclude.xlsx"
    wb = Workbook()
    sheet = wb.active.title
    styles = PoiExcelUtil.getCellStyles(wb)
    lstHeader = [["序号", "序号"], ["名称", "名称"]]
    PoiExcelUtil.writeHeaderByStyleExclude(wb, sheet, lstHeader, styles, 1, [1])
    PoiExcelUtil.writeCellDataByStyle(wb, sheet, styles, 2, 0, 1)
    PoiExcelUtil.writeCellDataByStyle(wb, sheet, styles, 2, 1, "商品")
    PoiExcelUtil.write(wb, str(out))

    data = PoiExcelUtil.toMapFull(str(out), sheet, 0, 0, -1, 2, -1, None)
    assert len(data) == 1
    assert data[0]["序号"] == "1"
    assert data[0]["名称"] == "商品"
